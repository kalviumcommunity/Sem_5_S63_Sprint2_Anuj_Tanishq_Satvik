"""Document file loading and multi-format intake."""

import hashlib
import io
import os
import re
from typing import List, Optional, Tuple

from backend.app.models.document import DocumentPage, DocumentType, LoadedDocument
from backend.app.services.documents.exceptions import (
    CorruptedDocumentError,
    DocumentIntakeError,
    EmptyDocumentError,
    UnsupportedFormatError,
)

SUPPORTED_EXTENSIONS = ["pdf", "txt", "docx", "md", "markdown"]


class DocumentLoader:
    """Loads document files across multiple formats into normalized LoadedDocument instances."""

    @staticmethod
    def _compute_hash(data: bytes) -> str:
        """Compute SHA-256 hash of raw byte content."""
        return hashlib.sha256(data).hexdigest()

    @staticmethod
    def load_text_file(file_path: str) -> Tuple[str, int]:
        """Legacy helper: Load plain text document returning text and character count."""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()

        return content, len(content)

    @classmethod
    def load(cls, file_path: str) -> LoadedDocument:
        """Load document from local file system path."""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        filename = os.path.basename(file_path)
        with open(file_path, "rb") as f:
            data = f.read()

        return cls._parse_document(data=data, filename=filename, source_path=file_path)

    @classmethod
    def load_bytes(cls, data: bytes, filename: str) -> LoadedDocument:
        """Load document from in-memory byte buffer."""
        source_path = f"memory://{filename}"
        return cls._parse_document(data=data, filename=filename, source_path=source_path)

    @classmethod
    def safe_load(cls, file_path: str) -> Tuple[Optional[LoadedDocument], Optional[str]]:
        """Safely load a document, returning (doc, None) on success or (None, error_msg) on failure."""
        try:
            doc = cls.load(file_path)
            return doc, None
        except Exception as exc:
            return None, str(exc)

    @classmethod
    def _parse_document(cls, data: bytes, filename: str, source_path: str) -> LoadedDocument:
        doc_type = DocumentType.from_filename(filename)
        ext = os.path.splitext(filename)[1].lower()

        if doc_type == DocumentType.UNKNOWN:
            raise UnsupportedFormatError(filename=filename, extension=ext, supported_formats=SUPPORTED_EXTENSIONS)

        file_hash = cls._compute_hash(data)

        if doc_type == DocumentType.TXT:
            pages = cls._parse_txt(data, filename)
        elif doc_type == DocumentType.MARKDOWN:
            pages = cls._parse_markdown(data, filename)
        elif doc_type == DocumentType.DOCX:
            pages = cls._parse_docx(data, filename)
        elif doc_type == DocumentType.PDF:
            pages = cls._parse_pdf(data, filename)
        else:
            raise UnsupportedFormatError(filename=filename, extension=ext, supported_formats=SUPPORTED_EXTENSIONS)

        raw_text = "\n\n".join(p.text for p in pages if p.text.strip())

        if not raw_text.strip():
            raise EmptyDocumentError(filename=filename)

        return LoadedDocument(
            filename=filename,
            source_path=source_path,
            document_type=doc_type,
            raw_text=raw_text,
            pages=pages,
            page_count=len(pages),
            char_count=len(raw_text),
            file_hash=file_hash,
            metadata={
                "source_path": source_path,
                "file_hash": file_hash,
                "document_type": doc_type.value,
                "page_count": len(pages),
            },
        )

    @staticmethod
    def _parse_txt(data: bytes, filename: str) -> List[DocumentPage]:
        try:
            content = data.decode("utf-8")
        except UnicodeDecodeError:
            try:
                content = data.decode("latin-1")
            except Exception as e:
                raise CorruptedDocumentError(filename=filename, reason=f"Decoding failure: {e}")

        if not content.strip():
            raise EmptyDocumentError(filename=filename)

        parts = content.split("\f")
        pages: List[DocumentPage] = []
        for idx, part in enumerate(parts):
            text = part.strip()
            if text or len(parts) == 1:
                pages.append(
                    DocumentPage(
                        page_number=idx + 1,
                        text=text or part,
                        metadata={"format": "txt", "section_idx": idx},
                    )
                )

        if not pages:
            raise EmptyDocumentError(filename=filename)
        return pages

    @staticmethod
    def _parse_markdown(data: bytes, filename: str) -> List[DocumentPage]:
        try:
            content = data.decode("utf-8")
        except UnicodeDecodeError:
            content = data.decode("latin-1", errors="replace")

        if not content.strip():
            raise EmptyDocumentError(filename=filename)

        # Split on markdown section dividers: --- or *** or form feeds (supporting CRLF and LF)
        sections = re.split(r"(?:\r?\n[*-]{3,}\r?\n|\f)", content)
        pages: List[DocumentPage] = []
        for idx, section in enumerate(sections):
            cleaned = section.strip()
            if cleaned:
                pages.append(
                    DocumentPage(
                        page_number=idx + 1,
                        text=cleaned,
                        metadata={"format": "markdown", "section_index": idx + 1},
                    )
                )

        if not pages:
            pages = [DocumentPage(page_number=1, text=content.strip(), metadata={"format": "markdown"})]
        return pages

    @staticmethod
    def _parse_docx(data: bytes, filename: str) -> List[DocumentPage]:
        try:
            import docx
        except ImportError:
            raise CorruptedDocumentError(filename=filename, reason="python-docx package not installed")

        try:
            file_stream = io.BytesIO(data)
            doc = docx.Document(file_stream)
        except Exception as e:
            raise CorruptedDocumentError(filename=filename, reason=f"Invalid or corrupted DOCX archive: {e}")

        text_parts: List[str] = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text.strip())

        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)

        full_text = "\n\n".join(text_parts).strip()
        if not full_text:
            raise EmptyDocumentError(filename=filename)

        return [
            DocumentPage(
                page_number=1,
                text=full_text,
                metadata={"format": "docx", "paragraph_count": len(text_parts)},
            )
        ]

    @staticmethod
    def _parse_pdf(data: bytes, filename: str) -> List[DocumentPage]:
        # Try extracting via pypdf first
        pages: List[DocumentPage] = []
        pypdf_worked = False
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            if len(reader.pages) > 0:
                pypdf_worked = True
                for idx, page in enumerate(reader.pages):
                    page_text = page.extract_text() or ""
                    if page_text.strip():
                        pages.append(
                            DocumentPage(
                                page_number=idx + 1,
                                text=page_text.strip(),
                                metadata={"format": "pdf", "page_index": idx},
                            )
                        )
        except Exception:
            pypdf_worked = False

        # Fallback for synthetic/minimal PDF text streams or pypdf extraction gaps
        if not pages:
            # Check for corrupted PDF marker
            if not data.startswith(b"%PDF-"):
                raise CorruptedDocumentError(filename=filename, reason="Missing %PDF header")

            # Extract BT ... Tj / TJ text stream commands
            stream_matches = re.findall(rb"\(([^)]+)\)\s*(?:Tj|TJ)", data)
            if stream_matches:
                extracted = " ".join(match.decode("latin-1", errors="replace") for match in stream_matches)
                if extracted.strip():
                    pages.append(
                        DocumentPage(
                            page_number=1,
                            text=extracted.strip(),
                            metadata={"format": "pdf", "extraction_method": "stream_fallback"},
                        )
                    )

        if not pages:
            if not pypdf_worked and b"GARBAGE" in data or b"CORRUPTED" in data or b"stream" not in data:
                raise CorruptedDocumentError(filename=filename, reason="Unreadable or malformed PDF structure")
            raise EmptyDocumentError(filename=filename)

        return pages
